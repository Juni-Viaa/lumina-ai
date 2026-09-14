from pathlib import Path
from tempfile import TemporaryDirectory
from unittest.mock import MagicMock, patch

from django.test import SimpleTestCase
from langchain_core.documents import Document

from .services import (
    _load_docx_hybrid,
    _merge_pdf_page_text,
    _pdf_pages_requiring_ocr,
)
from .vision_service import (
    VisualAnalysis,
    _parse_visual_analysis,
    format_visual_analysis,
)


class PdfHybridOcrTests(SimpleTestCase):
    def test_merge_replaces_empty_page_with_ocr(self):
        self.assertEqual(_merge_pdf_page_text("", "Teks hasil OCR"), "Teks hasil OCR")

    def test_merge_does_not_duplicate_same_text(self):
        self.assertEqual(
            _merge_pdf_page_text("Halo, dunia!", "Halo dunia"),
            "Halo, dunia!",
        )

    @patch("pypdf.PdfReader")
    def test_selects_short_and_image_bearing_pages(self, reader_mock):
        short_page = MagicMock()
        short_page.images = []
        image_page = MagicMock()
        image_page.images = [object()]
        text_page = MagicMock()
        text_page.images = []
        reader_mock.return_value.pages = [short_page, image_page, text_page]
        docs = [
            Document(page_content="sedikit"),
            Document(page_content="teks digital yang panjang " * 10),
            Document(page_content="teks digital yang panjang " * 10),
        ]

        self.assertEqual(_pdf_pages_requiring_ocr(MagicMock(), docs), [0, 1])


class DocxHybridOcrTests(SimpleTestCase):
    @patch("ingest.ocr_service.ocr_image")
    @patch("ingest.vision_service.analyze_image")
    def test_extracts_and_inserts_embedded_image_text(
        self,
        vision_mock,
        ocr_mock,
    ):
        from docx import Document as WordDocument
        from PIL import Image

        ocr_mock.return_value = [{"text": "standard minmax robust"}]
        vision_mock.return_value = None
        with TemporaryDirectory() as temp_dir:
            path = Path(temp_dir) / "example.docx"
            image_path = Path(temp_dir) / "scaler.png"
            Image.new("RGB", (200, 100), "white").save(image_path)
            word_document = WordDocument()
            word_document.add_paragraph("Normalisasi data")
            paragraph = word_document.add_paragraph("Pilihan scaler")
            paragraph.add_run().add_picture(str(image_path))
            word_document.save(path)

            docs = _load_docx_hybrid(path, None, None)

        self.assertEqual(len(docs), 2)
        self.assertIn("standard minmax robust", docs[1].page_content)
        self.assertTrue(docs[1].metadata["ocr_used"])
        ocr_mock.assert_called_once()


class VisionAnalysisTests(SimpleTestCase):
    def test_parses_and_formats_structured_visual_analysis(self):
        raw = """```json
        {
          "visual_type": "ui_screenshot",
          "title": "Normalisasi",
          "description": "Menu pengaturan scaler",
          "visible_text": ["standard", "minmax", "robust"],
          "key_facts": ["Terdapat tiga pilihan scaler"],
          "relationships": []
        }
        ```"""

        analysis = _parse_visual_analysis(raw)
        formatted = format_visual_analysis(analysis)

        self.assertIsInstance(analysis, VisualAnalysis)
        self.assertIn("standard; minmax; robust", formatted)
        self.assertIn("Terdapat tiga pilihan scaler", formatted)
