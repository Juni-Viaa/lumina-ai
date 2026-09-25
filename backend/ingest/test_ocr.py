from pathlib import Path
from io import BytesIO
from tempfile import TemporaryDirectory
from types import SimpleNamespace
from unittest.mock import MagicMock, patch

from django.test import SimpleTestCase
from django.core.files.uploadedfile import SimpleUploadedFile
from langchain_core.documents import Document

from .services import (
    _load_document,
    _load_docx_hybrid,
    _load_pdf_hybrid,
    _merge_pdf_page_text,
    _ocr_confidence,
    _pdf_pages_requiring_ocr,
    _preprocess_documents,
)
from .ocr_service import _analyze_pdf_image_regions, _poppler_path, ocr_pdf_pages
from .serializers import IngestUploadSerializer
from .vision_service import (
    VisualAnalysis,
    _parse_visual_analysis,
    analyze_image,
    format_visual_analysis,
)


class ImageUploadValidationTests(SimpleTestCase):
    def test_valid_png_remains_uploadable(self):
        from PIL import Image

        image_data = BytesIO()
        Image.new("RGB", (20, 20), "white").save(image_data, format="PNG")
        upload = SimpleUploadedFile("scan.png", image_data.getvalue(), content_type="image/png")
        serializer = IngestUploadSerializer(data={"document": upload})
        self.assertTrue(serializer.is_valid(), serializer.errors)
        self.assertEqual(upload.tell(), 0)

    def test_invalid_image_is_rejected(self):
        upload = SimpleUploadedFile("scan.png", b"not an image", content_type="image/png")
        serializer = IngestUploadSerializer(data={"document": upload})
        self.assertFalse(serializer.is_valid())


class PdfHybridOcrTests(SimpleTestCase):
    @patch("ingest.vision_service.analyze_image")
    @patch("ingest.ocr_service.ocr_image")
    def test_analyzes_embedded_pdf_image_instead_of_decorative_icon(self, ocr_mock, vision_mock):
        from PIL import Image

        def embedded_image(size: tuple[int, int]) -> SimpleNamespace:
            image = Image.new("RGB", size, "white")
            data = BytesIO()
            image.save(data, format="PNG")
            return SimpleNamespace(image=image, data=data.getvalue())

        chart = embedded_image((400, 250))
        page = SimpleNamespace(images=[embedded_image((20, 20)), chart, chart])
        ocr_mock.return_value = [{"text": "standard minmax robust", "confidence": 0.86}]
        vision_mock.return_value = VisualAnalysis(
            visual_type="ui_screenshot", description="Pilihan scaler pada gambar",
        )

        with TemporaryDirectory() as temp_dir:
            regions = _analyze_pdf_image_regions(page, 2, Path(temp_dir), "Normalisasi data")

        self.assertEqual(len(regions), 1)
        self.assertEqual(regions[0]["image_ref"], "page-2-image-2")
        self.assertEqual(regions[0]["confidence"], 0.86)
        self.assertEqual(regions[0]["text"], "standard minmax robust")
        self.assertEqual(regions[0]["visual_analysis"]["description"], "Pilihan scaler pada gambar")
        self.assertEqual(vision_mock.call_args.args[2], "Normalisasi data")

    @patch("ingest.ocr_service.config.VISION_PDF_MAX_IMAGES_PER_PAGE", 2)
    @patch("ingest.vision_service.analyze_image")
    @patch("ingest.ocr_service.ocr_image", return_value=[])
    def test_limits_vision_calls_per_pdf_page(self, _ocr_mock, vision_mock):
        from PIL import Image

        images = []
        for index in range(5):
            image = Image.new("RGB", (300, 200), (index * 30, 0, 0))
            data = BytesIO()
            image.save(data, format="PNG")
            images.append(SimpleNamespace(image=image, data=data.getvalue()))
        vision_mock.return_value = VisualAnalysis(
            visual_type="diagram", description="Diagram pada halaman PDF",
        )

        with TemporaryDirectory() as temp_dir:
            regions = _analyze_pdf_image_regions(
                SimpleNamespace(images=images), 1, Path(temp_dir), "",
            )

        self.assertEqual(len(regions), 2)
        self.assertEqual(vision_mock.call_count, 2)

    @patch("ingest.vision_service.analyze_image")
    @patch("ingest.ocr_service._analyze_pdf_image_regions")
    @patch("ingest.ocr_service.ocr_image")
    @patch("pypdf.PdfReader")
    @patch("pdf2image.convert_from_path")
    @patch("ingest.ocr_service.config.GEMINI_API_KEY", "test-key")
    @patch("ingest.ocr_service.config.VISION_ENABLED", True)
    def test_pdf_vision_uses_region_and_falls_back_for_scanned_page(
        self, convert_mock, reader_mock, ocr_mock, regions_mock, vision_mock,
    ):
        from PIL import Image

        convert_mock.return_value = [Image.new("RGB", (300, 300), "white")]
        reader_mock.return_value.pages = [object(), object()]
        ocr_mock.return_value = [{"text": "pilihan scaler", "confidence": 0.8}]
        regions_mock.side_effect = [
            [{"image_ref": "page-1-image-1", "text": "pilihan scaler", "confidence": 0.8,
              "visual_analysis": {"visual_type": "ui_screenshot", "description": "Menu scaler"}}],
            [],
        ]
        vision_mock.return_value = VisualAnalysis(visual_type="chart", description="Grafik hasil")

        results = ocr_pdf_pages(
            Path("example.pdf"), [0, 1], analyze_visuals=True,
            page_contexts={0: "Pengaturan normalisasi"},
        )

        self.assertEqual(len(results[0]["visual_regions"]), 1)
        self.assertIsNone(results[0]["visual_analysis"])
        self.assertIsNotNone(results[1]["visual_analysis"])
        self.assertEqual(vision_mock.call_count, 1)

    @patch("ingest.ocr_service.ocr_pdf_pages")
    @patch("ingest.services._pdf_pages_requiring_ocr", return_value=[0])
    @patch("ingest.services.PyPDFLoader")
    def test_pdf_regions_keep_distinct_source_references(self, loader_mock, _pages_mock, ocr_mock):
        loader_mock.return_value.load.return_value = [Document(
            page_content="Normalisasi menggunakan berbagai pilihan scaler.", metadata={"page": 0}
        )]
        ocr_mock.return_value = [{
            "page_num": 1, "text": "Teks halaman PDF", "confidence": 0.8,
            "visual_analysis": None,
            "visual_regions": [{
                "image_ref": "page-1-image-1", "text": "standard minmax robust", "confidence": 0.91,
                "visual_analysis": {"visual_type": "ui_screenshot", "description": "Menu scaler"},
            }],
        }]

        docs = _load_pdf_hybrid(Path("sample.pdf"), None, None)

        region_docs = [doc for doc in docs if doc.metadata.get("image_ref") == "page-1-image-1"]
        self.assertEqual([doc.metadata["source_type"] for doc in region_docs], ["ocr", "vision"])
        self.assertEqual(region_docs[0].metadata["ocr_confidence"], 0.91)
        self.assertEqual(region_docs[0].metadata["page"], 0)

    @patch("ingest.ocr_service.ocr_pdf_pages")
    @patch("ingest.services._pdf_pages_requiring_ocr", return_value=[0])
    @patch("ingest.services.PyPDFLoader")
    def test_keeps_pdf_ocr_and_vision_separate(self, loader_mock, _pages_mock, ocr_mock):
        loader_mock.return_value.load.return_value = [Document(
            page_content="Teks digital pada halaman pertama dokumen ujian.", metadata={"page": 0}
        )]
        ocr_mock.return_value = [{
            "page_num": 1,
            "text": "Teks hasil OCR pada gambar halaman pertama.",
            "confidence": 0.82,
            "visual_analysis": {
                "visual_type": "ui_screenshot",
                "description": "Tampilan soal ujian pada halaman pertama.",
                "visible_text": ["Soal ujian"],
                "key_facts": [],
                "relationships": [],
            },
        }]

        docs = _load_pdf_hybrid(Path("sample.pdf"), None, None)

        self.assertEqual([doc.metadata["source_type"] for doc in docs], [
            "document_text", "ocr", "vision",
        ])
        self.assertEqual(docs[1].metadata["ocr_confidence"], 0.82)
        self.assertEqual(docs[1].metadata["image_ref"], docs[2].metadata["image_ref"])
        self.assertNotIn("[Analisis visual]", docs[1].page_content)

    @patch("ingest.ocr_service.os.name", "posix")
    def test_linux_uses_poppler_from_path(self):
        self.assertIsNone(_poppler_path())

    def test_merge_replaces_empty_page_with_ocr(self):
        self.assertEqual(_merge_pdf_page_text("", "Teks hasil OCR"), "Teks hasil OCR")

    def test_merge_does_not_duplicate_same_text(self):
        self.assertEqual(
            _merge_pdf_page_text("Halo, dunia!", "Halo dunia"),
            "Halo, dunia!",
        )

    @patch("pypdf.PdfReader")
    def test_selects_short_and_image_bearing_pages(self, reader_mock):
        short_page = MagicMock()
        short_page.images = []
        image_page = MagicMock()
        image_page.images = [object()]
        text_page = MagicMock()
        text_page.images = []
        reader_mock.return_value.pages = [short_page, image_page, text_page]
        docs = [
            Document(page_content="sedikit"),
            Document(page_content="teks digital yang panjang " * 10),
            Document(page_content="teks digital yang panjang " * 10),
        ]

        self.assertEqual(_pdf_pages_requiring_ocr(MagicMock(), docs), [0, 1])


class DocxHybridOcrTests(SimpleTestCase):
    @patch("ingest.services.config.VISION_DOCX_MIN_IMAGE_PIXELS", 1)
    @patch("ingest.services.config.VISION_DOCX_MAX_IMAGES", 2)
    @patch("ingest.ocr_service.ocr_image")
    @patch("ingest.vision_service.analyze_image")
    def test_limits_docx_vision_to_largest_images(
        self, vision_mock, ocr_mock,
    ):
        from docx import Document as WordDocument
        from PIL import Image

        ocr_mock.return_value = [{"text": "teks gambar", "confidence": 0.9}]
        vision_mock.return_value = VisualAnalysis(
            visual_type="ui_screenshot", description="Tampilan aplikasi",
        )
        with TemporaryDirectory() as temp_dir:
            path = Path(temp_dir) / "images.docx"
            word_document = WordDocument()
            for index, size in enumerate(((80, 80), (200, 100), (300, 200)), start=1):
                image_path = Path(temp_dir) / f"image-{index}.png"
                Image.new("RGB", size, (index * 40, 0, 0)).save(image_path)
                word_document.add_paragraph(f"Gambar {index}").add_run().add_picture(
                    str(image_path)
                )
            word_document.save(path)

            docs = _load_docx_hybrid(path, None, None)

        self.assertEqual(ocr_mock.call_count, 3)
        self.assertEqual(vision_mock.call_count, 2)
        self.assertEqual(
            len([doc for doc in docs if doc.metadata.get("source_type") == "vision"]),
            2,
        )

    @patch("ingest.ocr_service.ocr_image")
    @patch("ingest.vision_service.analyze_image")
    def test_extracts_and_inserts_embedded_image_text(
        self,
        vision_mock,
        ocr_mock,
    ):
        from docx import Document as WordDocument
        from PIL import Image

        ocr_mock.return_value = [{"text": "standard minmax robust", "confidence": 0.9}]
        vision_mock.return_value = None
        with TemporaryDirectory() as temp_dir:
            path = Path(temp_dir) / "example.docx"
            image_path = Path(temp_dir) / "scaler.png"
            Image.new("RGB", (200, 100), "white").save(image_path)
            word_document = WordDocument()
            word_document.add_paragraph("Normalisasi data")
            paragraph = word_document.add_paragraph("Pilihan scaler")
            paragraph.add_run().add_picture(str(image_path))
            word_document.save(path)

            docs = _load_docx_hybrid(path, None, None)

        self.assertEqual(len(docs), 3)
        self.assertEqual(docs[2].page_content, "standard minmax robust")
        self.assertEqual(docs[2].metadata["source_type"], "ocr")
        self.assertEqual(docs[2].metadata["ocr_confidence"], 0.9)
        self.assertTrue(docs[2].metadata["ocr_used"])
        self.assertEqual(len(_preprocess_documents(docs, None, None)), 1)
        ocr_mock.assert_called_once()


class ImageProvenanceTests(SimpleTestCase):
    @patch("ingest.vision_service.analyze_image")
    @patch("ingest.ocr_service.ocr_document")
    def test_image_keeps_ocr_and_vision_separate(self, ocr_mock, vision_mock):
        ocr_mock.return_value = {
            "text": "standard minmax robust",
            "raw_results": [
                {"text": "standard", "confidence": 0.9},
                {"text": "minmax", "confidence": 0.8},
            ],
        }
        vision_mock.return_value = VisualAnalysis(
            visual_type="ui_screenshot",
            description="Menu normalisasi",
            visible_text=["standard", "minmax", "robust"],
        )

        docs = _load_document(Path("normalisasi.png"), None, None)

        self.assertEqual([doc.metadata["source_type"] for doc in docs], ["ocr", "vision"])
        self.assertEqual(docs[0].metadata["ocr_confidence"], 0.85)
        self.assertEqual(docs[0].metadata["image_ref"], "normalisasi.png")
        self.assertEqual(docs[0].metadata["image_ref"], docs[1].metadata["image_ref"])
        self.assertIn("standard minmax robust", docs[0].page_content)
        self.assertIn("[Analisis visual]", docs[1].page_content)

    def test_missing_ocr_scores_are_not_treated_as_zero(self):
        self.assertIsNone(_ocr_confidence([{"text": "hasil"}]))


class VisionAnalysisTests(SimpleTestCase):
    @patch("ingest.vision_service.config.GEMINI_API_KEY", "test-key")
    @patch("ingest.vision_service.config.VISION_ENABLED", True)
    @patch("ingest.vision_service.ChatGoogleGenerativeAI")
    def test_uses_structured_output_and_validates_result(self, model_class):
        model_class.return_value.with_structured_output.return_value.invoke.return_value = {
            "visual_type": "ui_screenshot", "description": "Menu scaler",
            "visible_text": ["standard", "minmax", "robust"],
        }
        with TemporaryDirectory() as temp_dir:
            path = Path(temp_dir) / "image.png"
            path.write_bytes(b"image")
            result = analyze_image(path)
        self.assertEqual(result.description, "Menu scaler")
        model_class.return_value.with_structured_output.assert_called_once_with(
            VisualAnalysis, method="json_schema"
        )

    @patch("ingest.vision_service.config.GEMINI_API_KEY", "test-key")
    @patch("ingest.vision_service.config.VISION_ENABLED", True)
    @patch("ingest.vision_service.ChatGoogleGenerativeAI")
    def test_invalid_structured_result_is_discarded(self, model_class):
        model_class.return_value.with_structured_output.return_value.invoke.return_value = {
            "visual_type": "made-up", "description": "",
        }
        with TemporaryDirectory() as temp_dir:
            path = Path(temp_dir) / "image.png"
            path.write_bytes(b"image")
            self.assertIsNone(analyze_image(path))

    def test_parses_and_formats_structured_visual_analysis(self):
        raw = """```json
        {
          "visual_type": "ui_screenshot",
          "title": "Normalisasi",
          "description": "Menu pengaturan scaler",
          "visible_text": ["standard", "minmax", "robust"],
          "key_facts": ["Terdapat tiga pilihan scaler"],
          "relationships": []
        }
        ```"""

        analysis = _parse_visual_analysis(raw)
        formatted = format_visual_analysis(analysis)

        self.assertIsInstance(analysis, VisualAnalysis)
        self.assertIn("standard; minmax; robust", formatted)
        self.assertIn("Terdapat tiga pilihan scaler", formatted)
